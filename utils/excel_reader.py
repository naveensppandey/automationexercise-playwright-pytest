import os
import openpyxl

def read_excel_cases(file_path="test_data/test_cases.xlsx", sheet_name="TestCases"):
    """
    Reads test cases from an Excel (.xlsx) file and returns them as a list of dictionaries.
    Uses openpyxl to inspect row by row.
    """
    # Locate file relative to project root
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    full_path = os.path.join(project_root, file_path)

    if not os.path.exists(full_path):
        raise FileNotFoundError(f"Excel file not found at: {full_path}")

    # Load workbook and select worksheet
    workbook = openpyxl.load_workbook(full_path)
    sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active

    # Read column header names from the first row
    headers = [cell.value for cell in sheet[1]]

    # Read all remaining rows
    cases = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(row):  # Skip completely empty rows
            case_data = dict(zip(headers, row))
            cases.append(case_data)

    workbook.close()
    return cases

def get_test_case_by_id(tc_id, file_path="test_data/test_cases.xlsx"):
    """Retrieve a single test case dictionary by its ID (e.g., 'TC01')."""
    cases = read_excel_cases(file_path)
    for case in cases:
        if case.get("Test Case ID") == tc_id:
            return case
    return None
