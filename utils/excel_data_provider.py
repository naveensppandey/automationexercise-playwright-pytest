import os
import openpyxl

def read_excel_data(filename="test_data/test_cases.xlsx", sheet_name="TestCases", limit=None):
    """
    Reads an Excel (.xlsx) sheet using openpyxl and returns rows as a list of dictionaries.
    Handles empty rows safely and validates file and sheet existence.
    """
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    filepath = os.path.join(project_root, filename)

    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Excel test case file not found at: {filepath}")

    workbook = openpyxl.load_workbook(filepath)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in workbook {filename}. Available: {workbook.sheetnames}")

    sheet = workbook[sheet_name]
    rows = []
    headers = [cell.value for cell in sheet[1]]

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(row):
            row_dict = dict(zip(headers, row))
            rows.append(row_dict)
            if limit and len(rows) >= limit:
                break

    workbook.close()
    return rows

def get_test_case_by_id(tc_id, filename="test_data/test_cases.xlsx", sheet_name="TestCases"):
    """Returns a single test case dictionary matching the specified Test Case ID."""
    all_cases = read_excel_data(filename=filename, sheet_name=sheet_name)
    for tc in all_cases:
        if tc.get("Test Case ID") == tc_id:
            return tc
    return None

