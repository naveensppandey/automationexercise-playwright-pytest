import os
import re
import pytest
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from playwright.sync_api import sync_playwright
from utils.excel_data_provider import read_excel_data

# Global dictionary to track execution results for all 50 Excel test cases
EXECUTION_RESULTS = {}

def pytest_configure(config):
    # Register custom tc_id marker
    config.addinivalue_line("markers", "tc_id(*ids): associate test with Excel test case ID(s)")

    # Initialize execution results from test_data/test_cases.xlsx
    global EXECUTION_RESULTS
    try:
        cases = read_excel_data("test_data/test_cases.xlsx")
        for tc in cases:
            tc_id = tc.get("Test Case ID")
            if tc_id:
                EXECUTION_RESULTS[tc_id] = {
                    "Test Case ID": tc_id,
                    "Module": tc.get("Module", ""),
                    "Test Case": tc.get("Test Case", ""),
                    "Test Steps": tc.get("Test Steps", ""),
                    "Expected Result": tc.get("Expected Result", ""),
                    "Execution Status": "NOT AUTOMATED",
                    "Actual Result": "Not executed",
                    "Failure Reason": "",
                    "Screenshot Path": ""
                }
    except Exception as e:
        print(f"[Warning] Could not initialize Excel cases: {e}")

def pytest_addoption(parser):
    # Command line options to run browser in headed mode and choose browser (chromium or firefox)
    parser.addoption("--headed", action="store_true", default=False, help="Run browser in headed mode")
    parser.addoption("--browser", action="store", default="chromium", help="Browser type: chromium or firefox")

@pytest.fixture
def page(request):
    """
    Pytest Fixture for browser setup and teardown.
    - Supports Chromium and Firefox browsers via --browser option.
    - Launches browser before test.
    - Passes 'page' to test.
    - Closes browser after test completes.
    """
    is_headed = request.config.getoption("--headed")
    browser_name = request.config.getoption("--browser").lower()
    
    with sync_playwright() as playwright:
        # Launch Chromium or Firefox browser based on option with anti-bot options
        if browser_name == "firefox":
            browser = playwright.firefox.launch(headless=not is_headed)
        else:
            browser = playwright.chromium.launch(
                headless=not is_headed,
                args=["--disable-blink-features=AutomationControlled", "--no-sandbox"]
            )

        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
            viewport={"width": 1280, "height": 720}
        )
        # Abort third-party ad and analytics requests that stall automationexercise.com navigation
        context.route("**/*google-analytics*", lambda route: route.abort())
        context.route("**/*googlesyndication*", lambda route: route.abort())
        context.route("**/*doubleclick*", lambda route: route.abort())
        context.route("**/*adservice*", lambda route: route.abort())
        context.route("**/*pagead*", lambda route: route.abort())
        context.route("**/*adsbygoogle*", lambda route: route.abort())

        page_instance = context.new_page()
        
        # Give page instance to the test
        yield page_instance
        
        # Close browser after test finishes
        context.close()
        browser.close()

@pytest.hookimpl(tryfirst=True, hookwrapper=True)
def pytest_runtest_makereport(item, call):
    """Pytest Hook: Automatically captures screenshot on failure and records Excel test status."""
    outcome = yield
    report = outcome.get_result()

    if report.when == "call":
        # Identify associated Excel Test Case IDs
        tc_ids = []
        tc_marker = item.get_closest_marker("tc_id")
        if tc_marker:
            tc_ids.extend(tc_marker.args)
        
        # Also check for TC pattern in test name (e.g. test_tc01_... or test_tc08_tc10)
        found_in_name = re.findall(r"tc\d+", item.name, re.IGNORECASE)
        for found in found_in_name:
            norm_id = found.upper()
            if norm_id not in tc_ids:
                tc_ids.append(norm_id)

        screenshot_path = ""
        if report.failed:
            page_instance = item.funcargs.get("page")
            if page_instance:
                os.makedirs("screenshots", exist_ok=True)
                screenshot_path = f"screenshots/{item.name}_failure.png"
                try:
                    page_instance.screenshot(path=screenshot_path)
                    print(f"\n[Screenshot Saved] {screenshot_path}")
                except Exception:
                    pass

        # Update status for each identified TC ID
        for tid in tc_ids:
            if tid in EXECUTION_RESULTS:
                if report.passed:
                    EXECUTION_RESULTS[tid]["Execution Status"] = "PASS"
                    EXECUTION_RESULTS[tid]["Actual Result"] = f"Passed: {EXECUTION_RESULTS[tid]['Expected Result']}"
                    EXECUTION_RESULTS[tid]["Failure Reason"] = ""
                    EXECUTION_RESULTS[tid]["Screenshot Path"] = ""
                else:
                    EXECUTION_RESULTS[tid]["Execution Status"] = "FAIL"
                    EXECUTION_RESULTS[tid]["Actual Result"] = "Test step or assertion failed during execution"
                    EXECUTION_RESULTS[tid]["Failure Reason"] = str(report.longrepr)
                    EXECUTION_RESULTS[tid]["Screenshot Path"] = screenshot_path

def pytest_sessionfinish(session, exitstatus):
    """Pytest Hook: Generates reports/execution_results.xlsx after test suite completion."""
    if getattr(session.config.option, "collectonly", False):
        return
    if not EXECUTION_RESULTS:
        return

    os.makedirs("reports", exist_ok=True)
    report_file = os.path.join("reports", "execution_results.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Execution Results"

    headers = [
        "Test Case ID", "Module", "Test Case", "Test Steps",
        "Expected Result", "Execution Status", "Actual Result",
        "Failure Reason", "Screenshot Path"
    ]
    ws.append(headers)

    # Styling
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    pass_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="155724")
    fail_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    fail_font = Font(name="Calibri", size=10, bold=True, color="721C24")
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_idx = 2
    passed_count = 0
    failed_count = 0
    not_auto_count = 0

    for tid in sorted(EXECUTION_RESULTS.keys(), key=lambda x: int(re.sub(r'\D', '', x)) if re.sub(r'\D', '', x) else 0):
        data = EXECUTION_RESULTS[tid]
        status = data["Execution Status"]
        if status == "PASS":
            passed_count += 1
        elif status == "FAIL":
            failed_count += 1
        else:
            not_auto_count += 1

        row_values = [
            data["Test Case ID"],
            data["Module"],
            data["Test Case"],
            data["Test Steps"],
            data["Expected Result"],
            status,
            data["Actual Result"],
            data["Failure Reason"][:300] if data["Failure Reason"] else "",
            data["Screenshot Path"]
        ]
        ws.append(row_values)

        status_cell = ws.cell(row=row_idx, column=6)
        if status == "PASS":
            status_cell.fill = pass_fill
            status_cell.font = pass_font
        elif status == "FAIL":
            status_cell.fill = fail_fill
            status_cell.font = fail_font

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = thin_border
            if col_idx in [1, 2, 6]:
                c.alignment = Alignment(horizontal="center", vertical="top")
            else:
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        row_idx += 1

    # Adjust column widths
    col_widths = [14, 14, 28, 32, 32, 18, 32, 35, 25]
    for idx, width in enumerate(col_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = width

    wb.save(report_file)
    print(f"\n==================================================")
    print(f"[REPORT] EXCEL TEST EXECUTION REPORT GENERATED")
    print(f"File: {report_file}")
    print(f"Total: {len(EXECUTION_RESULTS)} | Passed: {passed_count} | Failed: {failed_count} | Not Automated: {not_auto_count}")
    print(f"==================================================\n")


